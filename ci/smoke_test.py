"""Yig'ilgan .exe ni haqiqatan ishga tushirib tekshiradi (GitHub Actions ichida).

Ishlatish: python ci/smoke_test.py dist/AlifboKonverter.exe
"""
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import docx

# Windows konsoli (cp1252) o'zbekcha harflarni chiqara olmaydi: UTF-8 majburlaymiz
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

exe = Path(sys.argv[1]).resolve()
assert exe.exists(), f"Topilmadi: {exe}"

work = Path(tempfile.mkdtemp())
src = work / "Test.docx"
d = docx.Document()
d.add_paragraph("O'zbekiston shahri, to'g'ri, G'oya, CHIROYLI, ta'lim, mas'hul, tong")
d.save(str(src))

env = dict(os.environ, APPDATA=str(work / "appdata"), XDG_CONFIG_HOME=str(work / "appdata"))
r = subprocess.run([str(exe), str(src), "-o", str(work / "out")], env=env, timeout=180)
assert r.returncode == 0, f"exe xato bilan tugadi: {r.returncode}"

out = work / "out" / "Test_yangi_alifbo.docx"
assert out.exists(), "Natija fayli yaratilmadi"
text = "".join(p.text for p in docx.Document(str(out)).paragraphs)
expected = "Özbekiston şahri, töğri, Ğoya, ÇIROYLI, ta’lim, mas’hul, tong"
assert text == expected, f"Kutilgan: {expected!r}\nOlindi:   {text!r}"
with zipfile.ZipFile(out) as z:
    assert z.testzip() is None
print("SMOKE TEST OK:", ascii(text))
  
