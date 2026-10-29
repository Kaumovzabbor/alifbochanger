import hashlib
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest
from lxml import etree

from app.converter import ConvertOptions
from app.docx_processor import (DocxError, convert_docx, convert_file, preview_docx)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
NSDECL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
)
DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'

CT = DECL + (
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Default Extension="png" ContentType="image/png"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    '<Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>'
    '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
    '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
    '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
    '<Override PartName="/word/charts/chart1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>'
    '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
    '</Types>')
RELS = DECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
    '</Relationships>')
DOC_RELS = DECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/>'
    '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>'
    '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
    '<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>'
    '</Relationships>')

PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360000002000001e221bc330000000049454e44ae426082")


def run(text, rpr="", tag="t"):
    sp = ' xml:space="preserve"'
    return f"<w:r>{('<w:rPr>' + rpr + '</w:rPr>') if rpr else ''}<w:{tag}{sp}>{text}</w:{tag}></w:r>"


def para(*runs):
    return "<w:p>" + "".join(runs) + "</w:p>"


def wrap(tag, body):
    return DECL + f"<w:{tag} {NSDECL}>{body}</w:{tag}>"


def make_docx(path, body, header="", footer="", footnotes="", comments="", chart="", core="", rels=True):
    parts = {
        "[Content_Types].xml": CT,
        "_rels/.rels": RELS,
        "word/_rels/document.xml.rels": DOC_RELS,
        "word/document.xml": wrap("document", f"<w:body>{body}</w:body>"),
        "word/header1.xml": wrap("hdr", header or para(run("sarlavha"))),
        "word/footer1.xml": wrap("ftr", footer or para(run("pastki"))),
        "word/footnotes.xml": wrap("footnotes", footnotes or '<w:footnote w:id="1">' + para(run("izoh")) + "</w:footnote>"),
        "word/comments.xml": wrap("comments", comments or '<w:comment w:id="0" w:author="A">' + para(run("fikr")) + "</w:comment>"),
        "word/media/image1.png": PNG,
    }
    if chart:
        parts["word/charts/chart1.xml"] = chart
    if core:
        parts["docProps/core.xml"] = core
    # [Content_Types].xml ataylab oxirida: natija birinchi bo'lishi kerak
    order = [k for k in parts if k != "[Content_Types].xml"] + ["[Content_Types].xml"]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for k in order:
            z.writestr(k, parts[k])
    return path


def texts(path, part="word/document.xml", tag="t"):
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read(part))
    return ["".join(x.text or "" for x in p.iter(f"{{{W}}}{tag}")) for p in root.iter(f"{{{W}}}p")]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_part(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name)


@pytest.fixture
def tmp(tmp_path):
    return tmp_path


def test_word_split_into_three_runs_keeps_formats(tmp):
    body = para(
        run("bo", "<w:b/>"),
        run("’", '<w:color w:val="FF0000"/>'),
        run("limi", "<w:i/>"),
        run(" va ", ""),
        run("SH", "<w:b/>"),
        run("AHAR", '<w:color w:val="00FF00"/>'),
    )
    src = make_docx(tmp / "a.docx", body)
    dst = tmp / "out.docx"
    convert_docx(src, dst)
    assert texts(dst)[0] == "bölimi va ŞAHAR"
    with zipfile.ZipFile(dst) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    runs = root.findall(".//w:r", NS)
    assert len(runs) == 6
    assert runs[0].find("w:rPr/w:b", NS) is not None
    assert runs[1].find("w:rPr/w:color", NS).get(f"{{{W}}}val") == "FF0000"
    assert runs[2].find("w:rPr/w:i", NS) is not None
    assert runs[0].find("w:t", NS).text == "bö"
    assert (runs[1].find("w:t", NS).text or "") == ""
    assert runs[2].find("w:t", NS).text == "limi"
    assert runs[4].find("w:rPr/w:b", NS) is not None
    # bosh/oxirgi probel yo'qolmaydi
    assert runs[3].find("w:t", NS).text == " va "


def test_digraph_across_runs_and_tab_is_boundary(tmp):
    body = para(run("s"), run("h"), run("ahar "), run("c"), "<w:r><w:tab/></w:r>", run("h"), run("i"))
    src = make_docx(tmp / "a.docx", body)
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    assert texts(dst)[0] == "şahar c" + "hi"   # tab c va h orasida: ç bo'lmaydi


def test_textbox_choice_and_fallback(tmp):
    box = (
        "<w:r><mc:AlternateContent>"
        "<mc:Choice Requires=\"wps\"><w:drawing><wp:inline><wp:docPr id=\"1\" name=\"Rasm 1\" descr=\"shahar rasmi\" title=\"Chiroyli\"/>"
        "<a:graphic><a:graphicData><wps:wsp><wps:txbx><w:txbxContent>"
        + para(run("o'zbek ", ""), run("textbox", "")) + "</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>"
        + para(run("o'zbek ", ""), run("textbox", "")) + "</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r>")
    body = para(run("oldin "), box, run(" keyin g'oya"))
    src = make_docx(tmp / "a.docx", body)
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    t = texts(dst)
    assert t.count("özbek textbox") == 2
    xml_t = read_part(dst, "word/document.xml").decode()
    assert "<w:t xml:space=\"preserve\">oldin </w:t>" in xml_t and " keyin ğoya" in xml_t
    xml = read_part(dst, "word/document.xml").decode()
    assert 'descr="şahar rasmi"' in xml and 'title="Çiroyli"' in xml
    assert 'name="Rasm 1"' in xml


def test_header_footer_footnote_comment(tmp):
    src = make_docx(
        tmp / "a.docx", para(run("x")),
        header=para(run("O'zbekiston ", ""), run("sh", ""), run("ahri", "")),
        footer=para(run("CHIROYLI")),
        footnotes='<w:footnote w:id="1">' + para(run("to'g'ri")) + "</w:footnote>",
        comments='<w:comment w:id="0" w:author="A">' + para(run("bo'lim")) + "</w:comment>")
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    assert texts(dst, "word/header1.xml") == ["Özbekiston şahri"]
    assert texts(dst, "word/footer1.xml") == ["ÇIROYLI"]
    assert texts(dst, "word/footnotes.xml") == ["töğri"]
    assert texts(dst, "word/comments.xml") == ["bölim"]


def test_tracked_changes_ins_del(tmp):
    body = para(
        "<w:ins w:id=\"1\" w:author=\"A\" w:date=\"2024-01-01T00:00:00Z\">" + run("yangi o'zbek") + "</w:ins>",
        "<w:del w:id=\"2\" w:author=\"A\" w:date=\"2024-01-01T00:00:00Z\">" + run(" eski shahar", tag="delText") + "</w:del>")
    src = make_docx(tmp / "a.docx", body)
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    assert texts(dst)[0] == "yangi özbek"
    assert texts(dst, tag="delText")[0] == " eski şahar"


def test_field_codes_and_hyperlink_untouched(tmp):
    body = para(
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>',
        '<w:r><w:instrText xml:space="preserve"> HYPERLINK "https://shahar.uz/chiroyli" </w:instrText></w:r>',
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>',
        run("shahar.uz"),
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>',
        run(" havola shahar"))
    src = make_docx(tmp / "a.docx", body)
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    xml = read_part(dst, "word/document.xml").decode()
    assert 'HYPERLINK "https://shahar.uz/chiroyli"' in xml
    assert texts(dst)[0] == "shahar.uz havola şahar"


def test_language_filter(tmp):
    body = para(
        run("church ", '<w:lang w:val="en-US"/>'),
        run("shahar ", '<w:lang w:val="uz-Latn-UZ"/>'),
        run("chiroyli ", '<w:lang w:val="ru-RU"/>'),
        run("ish"))
    src = make_docx(tmp / "a.docx", body)
    dst = tmp / "o.docx"
    st = convert_docx(src, dst, ConvertOptions(only_uzbek_lang=True))
    assert texts(dst)[0] == "church şahar chiroyli iş"
    assert st.skipped_lang_runs == 2
    dst2 = tmp / "o2.docx"
    convert_docx(src, dst2, ConvertOptions(only_uzbek_lang=False))
    assert texts(dst2)[0] == "çurç şahar çiroyli iş"


def test_chart_and_core_props(tmp):
    chart = DECL + (
        f'<c:chartSpace {NSDECL}><c:chart><c:title><c:tx><c:rich><a:p><a:r><a:t>Shahar bo\'yicha</a:t></a:r></a:p></c:rich></c:tx></c:title>'
        '<c:plotArea><c:ser><c:cat><c:strRef><c:strCache><c:pt idx="0"><c:v>Chilonzor</c:v></c:pt></c:strCache></c:strRef></c:cat>'
        '<c:val><c:numRef><c:numCache><c:pt idx="0"><c:v>12.5</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser></c:plotArea></c:chart></c:chartSpace>')
    core = DECL + (
        f'<cp:coreProperties {NSDECL}><dc:title>O\'zbek shahri</dc:title><dc:subject>bo\'lim</dc:subject>'
        '<cp:keywords>chiroyli, g\'oya</cp:keywords><dc:creator>Shuhrat Chorievich</dc:creator></cp:coreProperties>')
    src = make_docx(tmp / "a.docx", para(run("x")), chart=chart, core=core)
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    c = read_part(dst, "word/charts/chart1.xml").decode()
    assert "Çilonzor" in c and "12.5" in c and "Şahar böyiça" in c
    k = read_part(dst, "docProps/core.xml").decode()
    assert "Özbek şahri" in k and "<dc:subject>bölim</dc:subject>" in k
    assert "çiroyli, ğoya" in k
    assert "Shuhrat Chorievich" in k      # muallif ismi o'zgarmaydi


def test_original_untouched_and_other_parts_identical(tmp):
    src = make_docx(tmp / "a.docx", para(run("o'zbek shahar")))
    h = sha(src)
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    assert sha(src) == h
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(dst) as b:
        assert b.namelist()[0] == "[Content_Types].xml"
        assert sorted(a.namelist()) == sorted(b.namelist())
        for n in a.namelist():
            if n != "word/document.xml":
                assert a.read(n) == b.read(n), n
        assert a.read("word/media/image1.png") == b.read("word/media/image1.png") == PNG
    # XML deklaratsiyasi saqlanadi
    assert read_part(dst, "word/document.xml").startswith(DECL.encode())


def test_unchanged_parts_are_not_rewritten(tmp):
    src = make_docx(tmp / "a.docx", para(run("salom dunyo")))
    dst = tmp / "o.docx"
    st = convert_docx(src, dst)
    assert st.parts_changed == 0 and st.total == 0
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(dst) as b:
        for n in a.namelist():
            assert a.read(n) == b.read(n)


def test_idempotent_docx(tmp):
    src = make_docx(tmp / "a.docx", para(run("bo"), run("'limi shahar"), run(" g'oya")))
    d1, d2 = tmp / "d1.docx", tmp / "d2.docx"
    convert_docx(src, d1)
    st = convert_docx(d1, d2)
    assert st.total == 0
    assert read_part(d1, "word/document.xml") == read_part(d2, "word/document.xml")


def test_outer_paragraph_not_mixed_with_textbox_text(tmp):
    box = ("<w:r><w:drawing><wp:inline><a:graphic><a:graphicData><wps:wsp><wps:txbx><w:txbxContent>"
           + para(run("h")) + "</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>")
    body = para(run("c"), box, run("hirt"))
    src = make_docx(tmp / "a.docx", body)
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    # "c" + (textbox) + "hirt": chegara bor, "ch" birikmasi hosil bo'lmaydi
    assert "ç" not in "".join(texts(dst)) and texts(dst)[0] == "chhirt"


def test_preview_does_not_modify(tmp):
    src = make_docx(tmp / "a.docx", para(run("bo'lim"), run(" shahar"), run(" bo'lim")))
    h = sha(src)
    words, st = preview_docx(src)
    assert sha(src) == h
    assert words[("bo'lim", "bölim")] == 2
    assert words[("shahar", "şahar")] == 1


def test_xml_space_preserve_added(tmp):
    body = "<w:p><w:r><w:t>bo</w:t></w:r><w:r><w:t>'</w:t></w:r><w:r><w:t> lim shahar </w:t></w:r></w:p>"
    src = make_docx(tmp / "a.docx", body)
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    with zipfile.ZipFile(dst) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    last = root.findall(".//w:t", NS)[2]
    assert last.text == " lim şahar "
    assert last.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"


def test_stoplist_and_exceptions_in_docx(tmp):
    src = make_docx(tmp / "a.docx", para(run("Che"), run("lsea shahar mo'tabar")))
    dst = tmp / "o.docx"
    opt = ConvertOptions(stop_words=["Chelsea"], exceptions={"mo'tabar": "mo’tabar"})
    convert_docx(src, dst, opt)
    assert texts(dst)[0] == "Chelsea şahar mo’tabar"


# --- Xatolar -------------------------------------------------------------------------

def test_corrupt_zip(tmp):
    p = tmp / "bad.docx"
    p.write_bytes(b"PK\x03\x04 this is not a real zip" + b"\x00" * 50)
    with pytest.raises(DocxError) as e:
        convert_docx(p, tmp / "o.docx")
    assert "buzilgan" in str(e.value)


def test_not_a_word_zip(tmp):
    p = tmp / "x.docx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("hello.txt", "salom")
    with pytest.raises(DocxError):
        convert_docx(p, tmp / "o.docx")


def test_password_protected(tmp):
    p = tmp / "locked.docx"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100 + "EncryptedPackage".encode("utf-16-le") + b"\x00" * 100)
    with pytest.raises(DocxError) as e:
        convert_docx(p, tmp / "o.docx")
    assert "parol" in str(e.value)


def test_old_doc_format(tmp):
    p = tmp / "old.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 500)
    with pytest.raises(DocxError) as e:
        convert_docx(p, tmp / "o.docx")
    assert ".doc" in str(e.value)


def test_missing_file(tmp):
    with pytest.raises(DocxError):
        convert_docx(tmp / "yoq.docx", tmp / "o.docx")


def test_broken_xml_part(tmp):
    src = make_docx(tmp / "a.docx", para(run("x")))
    bad = tmp / "b.docx"
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(bad, "w") as b:
        for n in a.namelist():
            b.writestr(n, b"<w:document><w:p>" if n == "word/document.xml" else a.read(n))
    with pytest.raises(DocxError) as e:
        convert_docx(bad, tmp / "o.docx")
    assert "document.xml" in str(e.value)


# --- convert_file: nomlash, .bak ---------------------------------------------------

def test_convert_file_naming_and_unique(tmp):
    src = make_docx(tmp / "Hujjat.docx", para(run("shahar")))
    out1, _ = convert_file(src)
    out2, _ = convert_file(src)
    assert out1.name == "Hujjat_yangi_alifbo.docx"
    assert out2.name == "Hujjat_yangi_alifbo_2.docx"
    od = tmp / "natija"
    out3, _ = convert_file(src, output_dir=od)
    assert out3.parent == od and out3.exists()
    assert not list(tmp.glob(".alifbo_*"))     # vaqtinchalik fayllar qolmaydi


def test_convert_file_replace_with_bak(tmp):
    src = make_docx(tmp / "H.docx", para(run("shahar")))
    orig = src.read_bytes()
    out, _ = convert_file(src, replace_original=True)
    assert out == src
    assert (tmp / "H.docx.bak").read_bytes() == orig
    assert texts(src)[0] == "şahar"


def test_failed_conversion_leaves_no_files(tmp):
    p = tmp / "bad.docx"
    p.write_bytes(b"not zip")
    with pytest.raises(DocxError):
        convert_file(p)
    assert [f.name for f in tmp.iterdir()] == ["bad.docx"]


# --- Word/LibreOffice'da ochilish ---------------------------------------------------

def test_opens_with_python_docx(tmp):
    import docx
    src = make_docx(tmp / "a.docx", para(run("o'zbek shahar")))
    dst = tmp / "o.docx"
    convert_docx(src, dst)
    d = docx.Document(str(dst))
    assert d.paragraphs[0].text == "özbek şahar"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice yo'q")
def test_libreoffice_opens_and_renders(tmp):
    import docx
    d = docx.Document()
    d.add_heading("O'zbekiston Respublikasi", 1)
    p = d.add_paragraph()
    p.add_run("Maktabgacha va maktab ta'limi ").bold = True
    p.add_run("boshqarmasi, shahar, to'g'ri.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Chilonzor"
    t.rows[0].cells[1].text = "G'oya"
    d.sections[0].header.paragraphs[0].text = "YIG'MA JILD"
    src = tmp / "real.docx"
    d.save(str(src))
    dst = tmp / "real_out.docx"
    convert_docx(src, dst)
    outdir = tmp / "pdf"
    r = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(dst)],
                       capture_output=True, timeout=180)
    pdf = outdir / "real_out.pdf"
    assert pdf.exists(), r.stderr
    if shutil.which("pdftotext"):
        txt = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True).stdout
        for w in ("Özbekiston", "ta’limi", "boşqarmasi", "şahar", "töğri", "Çilonzor", "Ğoya", "YIĞMA"):
            assert w in txt, w
